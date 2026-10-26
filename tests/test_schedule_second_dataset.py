"""Tests for the second schedule-truth dataset (#1296): a synthetic office.

The schedule-truth benchmark (#1240) had one dataset with truth
annotations (``bsi-clinic``), so its numbers described one building. This
suite pins the second dataset: the generator is deterministic, its truth
freezes and refuses drift like the clinic's, and the bench scores it as
its own dataset (reported separately, not folded into an aggregate).
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts import schedule_bench  # noqa: E402
from scripts.make_synthetic_schedule_workbook import (  # noqa: E402
    SEED,
    build_rows,
    write_workbook,
)
from scripts.schedule_bench import load_config, render, run_schedule_bench  # noqa: E402
from scripts.schedule_truth import (  # noqa: E402
    SCHEMA,
    TruthMismatch,
    build_truth,
    load_or_freeze_truth,
)

DATASET = "synth-office"


def _workbook(tmp_path: Path, name: str = "synth-office.xlsx") -> Path:
    return write_workbook(tmp_path / name, seed=SEED, force=True)


def test_generator_is_deterministic_byte_for_byte(tmp_path: Path) -> None:
    """Same seed, same workbook bytes: the truth stays pinned to its source."""
    a = _workbook(tmp_path / "a")
    b = _workbook(tmp_path / "b")
    assert a.read_bytes() == b.read_bytes()
    import hashlib

    assert hashlib.sha256(a.read_bytes()).hexdigest() == hashlib.sha256(b.read_bytes()).hexdigest()


def test_generator_refuses_to_clobber_without_force(tmp_path: Path) -> None:
    out = tmp_path / "w.xlsx"
    write_workbook(out, seed=SEED)
    with pytest.raises(FileExistsError):
        write_workbook(out, seed=SEED)
    # --force overwrites deliberately.
    write_workbook(out, seed=SEED, force=True)


def test_generator_output_carries_the_four_stated_sheets(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(_workbook(tmp_path), read_only=True)
    try:
        assert {"Component", "Type", "Space", "Attribute"} <= set(wb.sheetnames)
        comp = list(wb["Component"].iter_rows(values_only=True))
        header = [str(h) for h in comp[0]]
        assert "Name" in header and "TypeName" in header and "Space" in header
        assert "TagNumber" in header
    finally:
        wb.close()


def test_generator_states_capacities_and_airflow_requirements(tmp_path: Path) -> None:
    """The metrics need something stated to measure."""
    truth = build_truth(_workbook(tmp_path))
    # Capacities live in the type descriptions, as stated text.
    assert any("kW" in t["description"] for t in truth["types"].values())
    # Airflow requirements live on the Attribute sheet, per component.
    assert truth["requirements"]
    units = {r["unit"] for reqs in truth["requirements"].values() for r in reqs}
    assert units & {"L/s", "CFM"}
    assert truth["spaces"]


def test_generator_records_the_untagged_row_honestly(tmp_path: Path) -> None:
    """No tag is invented: the row is recorded as untagged, with a reason."""
    truth = build_truth(_workbook(tmp_path))
    assert truth["counts"]["untagged"] == 1
    assert truth["untagged"][0]["reason"]
    assert "untagged" not in truth["rows"]


def test_truth_freezes_then_loads_and_refuses_drift(tmp_path: Path) -> None:
    wb_path = _workbook(tmp_path)
    truth, how = load_or_freeze_truth(wb_path)
    assert how == "frozen"
    assert truth["schema"] == SCHEMA
    assert (wb_path.parent / "truth.json").exists()
    truth2, how2 = load_or_freeze_truth(wb_path)
    assert how2 == "loaded"
    assert truth2["workbook_sha256"] == truth["workbook_sha256"]
    # A changed workbook drifts the frozen truth away: refuse, never re-read.
    wb_path.write_bytes(wb_path.read_bytes() + b"\x00")
    with pytest.raises(TruthMismatch):
        load_or_freeze_truth(wb_path)
    truth3, how3 = load_or_freeze_truth(wb_path, force=True)
    assert how3 == "refrozen"


def test_bench_scores_the_second_dataset(tmp_path: Path) -> None:
    from mech_schedule import parse_cobie_workbook

    wb_path = _workbook(tmp_path)
    truth = build_truth(wb_path)
    result = parse_cobie_workbook(wb_path)
    m = run_schedule_bench(truth, result)
    assert m["counts"]["truth_rows"] == truth["counts"]["tagged"]
    assert m["counts"]["matched_tags"] >= 1
    assert m["counts"]["extra_tags"] == 0
    assert 0.0 < m["tags"]["recall"] <= 1.0
    # The dataset has stated capacities and airflow: both metrics measure
    # something, rather than reporting an empty column.
    assert m["capacity"]["stated"] > 0
    assert m["airflow"]["stated"] > 0


def test_bench_reports_each_dataset_separately(tmp_path: Path) -> None:
    """The scoreboard has a row per dataset, not one aggregate number."""
    wb_path = _workbook(tmp_path)
    truth = build_truth(wb_path)
    from mech_schedule import parse_cobie_workbook

    metrics = run_schedule_bench(truth, parse_cobie_workbook(wb_path))
    metrics["truth"] = "frozen"
    rows = [
        schedule_bench.BenchRow(name="bsi-clinic", status="missing", detail="no workbook"),
        schedule_bench.BenchRow(name=DATASET, status="ok", extras=metrics),
    ]
    text = render(rows)
    assert "## Datasets" in text
    assert "| bsi-clinic |" in text
    assert f"| {DATASET} |" in text
    assert f"## {DATASET} (ok)" in text
    assert "## bsi-clinic (missing)" in text


def test_bench_config_registers_the_second_dataset() -> None:
    """The dataset is in the bench config, with its schedule workbook."""
    buildings = {b.name: b for b in load_config(ROOT / "scripts" / "benchmarks.json")}
    assert DATASET in buildings, sorted(buildings)
    workbook = buildings[DATASET].schedule_workbook
    assert workbook is not None
    assert workbook.name == "synth-office.xlsx"
    # The config path follows the established external-dataset convention:
    # a ``~/workspace/datasets/...`` path, expanded by the loader.
    raw = json.loads((ROOT / "scripts" / "benchmarks.json").read_text())
    entry = next(b for b in raw["buildings"] if b["name"] == DATASET)
    assert entry["schedule_workbook"].startswith("~/workspace/datasets/")
    assert "matchline-bench" in entry["schedule_workbook"]
    assert workbook.is_absolute()


def test_bench_scores_only_buildings_with_a_schedule_workbook(tmp_path: Path) -> None:
    """Datasets are the buildings that carry a workbook; the rest are counted."""
    config = tmp_path / "benchmarks.json"
    config.write_text(
        json.dumps(
            {
                "buildings": [
                    {"name": "with-workbook", "schedule_workbook": "w.xlsx"},
                    {"name": "no-workbook", "arch_ifc": "a.ifc"},
                ]
            }
        )
    )
    buildings = load_config(config)
    assert [b.name for b in buildings] == ["with-workbook", "no-workbook"]
    datasets = [b for b in buildings if b.schedule_workbook is not None]
    assert [b.name for b in datasets] == ["with-workbook"]


def test_generator_rows_are_reproducible_from_the_seed(tmp_path: Path) -> None:
    """The seed alone determines the sheet bodies (no wall-clock state)."""
    a = build_rows(SEED)
    b = build_rows(SEED)
    assert a == b
    c = build_rows(SEED + 1)
    assert c != a
