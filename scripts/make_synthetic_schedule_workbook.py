#!/usr/bin/env python3
"""Generate a second schedule-truth benchmark dataset (#1296).

The schedule-truth benchmark (#1240) had exactly one dataset with truth
annotations (``bsi-clinic``), so its numbers described one building and
could not generalize. This script writes a **second** dataset: a small
synthetic office building whose COBie-style workbook carries stated
equipment, spaces, capacities and airflow requirements, so the benchmark
has a second, deliberately different building to score against.

Provenance: **generated, not a real building.** Every number, tag and room
comes from this generator; nothing here is copied from, or derived from, a
real project. The generator is deterministic: the same seed writes the
same workbook bytes (same sha256), so the frozen truth stays pinned to
the workbook it was read from. It follows the spirit of the #1241
scan-degradation family: a seeded, reproducible synthetic source whose
truth is generated, never hand-edited.

The workbook is written with openpyxl only (the same reader
``scripts/schedule_truth.py`` uses) and carries the four sheets that
reader expects:

- ``Component``: one row per scheduled item (name, type, space, stated
  description, stated tag number or a tag embedded in the name);
- ``Type``: the type catalogue, where the stated capacities live;
- ``Space``: the room list the service areas reference;
- ``Attribute``: the stated per-component airflow requirements
  (``Total`` / ``Outside Air``, L/s and CFM).

Nothing about the extractor's behaviour is encoded here: the workbook is
the *source*, and the truth is read back out of it independently by
``scripts/schedule_truth.py``.

Usage::

    python scripts/make_synthetic_schedule_workbook.py \
        --out ~/workspace/datasets/matchline-bench/synth-office/synth-office.xlsx \
        [--seed 1296] [--force]

The workbook is a generated output and is **not committed** (see
AGENTS.md: no datasets in code). The committed artefact is this script.
"""

from __future__ import annotations

import argparse
import io
import random
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

# A fixed creation stamp: the workbook must not carry wall-clock state,
# or two runs of the generator would not produce the same bytes.
FIXED_STAMP = datetime(2026, 1, 1, 0, 0, 0)
# Zip entries carry their own timestamps; fix them for byte determinism.
FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)

SEED = 1296

COMPONENT_HEADER = [
    "Name",
    "CreatedBy",
    "CreatedOn",
    "Category",
    "System",
    "TypeName",
    "Space",
    "ExtObject",
    "ExtIdentifier",
    "ExternalObject",
    "ExternalIdentifier",
    "Description",
    "TagNumber",
]
TYPE_HEADER = [
    "Name",
    "CreatedBy",
    "CreatedOn",
    "Category",
    "ExtObject",
    "ExtIdentifier",
    "Description",
]
SPACE_HEADER = [
    "Name",
    "CreatedBy",
    "CreatedOn",
    "Category",
    "FloorName",
    "ExtObject",
    "ExtIdentifier",
    "Description",
]
ATTRIBUTE_HEADER = [
    "Name",
    "CreatedBy",
    "CreatedOn",
    "Category",
    "SheetName",
    "RowName",
    "Value",
    "Unit",
]

# The building: a small two-storey office. Deliberately unlike the
# clinic (no clinical equipment, no water-cooled plant), so the bench
# scores two different building characters.
SPACES = [
    ("1.01", "Level 1", "Open Plan Office"),
    ("1.02", "Level 1", "Meeting Room"),
    ("1.03", "Level 1", "Manager Office"),
    ("1.04", "Level 1", "Kitchenette"),
    ("2.01", "Level 2", "Open Plan Office"),
    ("2.02", "Level 2", "Meeting Room"),
    ("2.03", "Level 2", "Server Room"),
    ("Roof Plant", "Roof", "Plant Deck"),
]

# Equipment catalogue. ``tag`` is embedded in the component name (the
# clinic convention: ``TagNumber`` reads ``n/a``), ``capacity`` is the
# stated type description, ``airflow`` the stated Attribute-sheet
# requirement. ``space`` is the service area the row states.
EQUIPMENT = [
    {
        "type": "Air handling unit",
        "tag": "AHU-1",
        "capacity": "Air handling unit, cooling capacity 45 kW",
        "space": "Roof Plant",
        "airflow": (("Total", "8000", "CFM"), ("Outside Air", "1200", "CFM")),
    },
    {
        "type": "Energy recovery ventilator",
        "tag": "ERV-1",
        "capacity": "Energy recovery ventilator, cooling capacity 12 kW",
        "space": "Roof Plant",
        "airflow": (("Total", "2500", "CFM"),),
    },
    {
        "type": "Variable air volume box",
        "tag": "VAV-1",
        "capacity": "Variable air volume box, cooling capacity 6 kW",
        "space": "1.01",
        "airflow": (("Total", "450", "CFM"), ("Outside Air", "90", "L/s")),
    },
    {
        "type": "Variable air volume box",
        "tag": "VAV-2",
        "capacity": "Variable air volume box, cooling capacity 6 kW",
        "space": "1.02",
        "airflow": (("Total", "380", "CFM"), ("Outside Air", "75", "L/s")),
    },
    {
        "type": "Variable air volume box",
        "tag": "VAV-3",
        "capacity": "Variable air volume box, cooling capacity 6 kW",
        "space": "1.03",
        "airflow": (("Total", "520", "CFM"), ("Outside Air", "100", "L/s")),
    },
    {
        "type": "Variable air volume box",
        "tag": "VAV-4",
        "capacity": "Variable air volume box, cooling capacity 6 kW",
        "space": "2.01",
        "airflow": (("Total", "520", "CFM"), ("Outside Air", "100", "L/s")),
    },
    {
        "type": "Variable air volume box",
        "tag": "VAV-5",
        "capacity": "Variable air volume box, cooling capacity 6 kW",
        "space": "2.02",
        "airflow": (("Total", "300", "CFM"),),
    },
    {
        "type": "Fan coil unit",
        "tag": "FCU-1",
        "capacity": "Fan coil unit, cooling capacity 6 kW",
        "space": "1.04",
        "airflow": (("Total", "220", "L/s"),),
    },
    {
        "type": "Fan coil unit",
        "tag": "FCU-2",
        "capacity": "Fan coil unit, cooling capacity 6 kW",
        "space": "2.03",
        "airflow": (("Total", "210", "L/s"),),
    },
    {
        "type": "Rooftop unit",
        "tag": "RTU-1",
        "capacity": "Rooftop unit, cooling capacity 60 kW",
        "space": "Roof Plant",
        "airflow": (("Total", "8000", "CFM"), ("Outside Air", "1500", "CFM")),
    },
    {
        "type": "Heat pump",
        "tag": "HP-1",
        "capacity": "Heat pump, heating capacity 16 kW",
        "space": "1.03",
        "airflow": (("Total", "150", "L/s"),),
    },
    {
        "type": "Exhaust fan",
        "tag": "EF-1",
        "capacity": "Exhaust fan, motor power 2 kW",
        "space": "1.04",
        "airflow": (("Total", "350", "CFM"),),
    },
    {
        "type": "Exhaust fan",
        "tag": "EF-2",
        "capacity": "Exhaust fan, motor power 2 kW",
        "space": "2.03",
        "airflow": (("Total", "280", "CFM"),),
    },
    {
        "type": "Chilled water pump",
        "tag": "P-1",
        "capacity": "Chilled water pump, rated motor power 22 kW",
        "space": "Roof Plant",
        "airflow": (),
    },
    {
        "type": "Air cooled chiller",
        "tag": "CH-1",
        "capacity": "Air cooled chiller, cooling capacity 35 ton",
        "space": "Roof Plant",
        "airflow": (),
    },
    {
        # An older workbook unit the extractor's capacity table does not
        # know: the bench then measures a real disagreement instead of a
        # tidy score.
        "type": "Split system condensing unit",
        "tag": "ACCU-1",
        "capacity": "Split system condensing unit, cooling capacity 12000 kJ/h",
        "space": "2.03",
        "airflow": (("Total", "400", "CFM"),),
    },
]

# Non-air-system rows: a schedule workbook also carries electrical and
# fire-safety items. They are stated rows, so the truth reader records
# them; the extractor excludes them by design, and the benchmark then
# shows the exclusion instead of hiding it.
NON_AIR_ROWS = [
    {
        "type": "Duplex receptacle",
        "tag": "RCP-1",
        "capacity": "Duplex receptacle",
        "space": "1.01",
    },
    {
        "type": "Fire alarm device",
        "tag": "FA-1",
        "capacity": "Fire alarm device",
        "space": "2.01",
    },
]

# One row with no stated tag at all: the truth reader records it under
# ``untagged`` with a reason, never drops it and never invents a tag.
UNTAGGED_ROW = {
    "type": "Motor control centre",
    "capacity": "Motor control centre",
    "space": "Roof Plant",
}


def _row(name: str, type_name: str, space: str, description: str, tag: str) -> list:
    """One Component row in the clinic's column order."""
    return [
        name,
        "Synth Schedule Generator",
        FIXED_STAMP,
        "Requirement",
        "",
        type_name,
        space,
        "",
        "",
        "",
        "",
        description,
        tag,
    ]


def _component_name(type_name: str, tag: str) -> str:
    """The component name carries the tag, as the clinic workbook does."""
    return f"{type_name.title()} {tag}" if tag else type_name.title()


def build_rows(seed: int) -> tuple[list[list], list[list], list[list], list[list]]:
    """Build the four sheet bodies from the seed (deterministically)."""
    rng = random.Random(seed)
    component_rows: list[list] = []
    type_rows: list[list] = []
    space_rows: list[list] = []
    attribute_rows: list[list] = []

    for number, floor, description in SPACES:
        space_rows.append(
            [
                number,
                "Synth Schedule Generator",
                FIXED_STAMP,
                "Requirement",
                floor,
                "",
                "",
                description,
            ]
        )

    seen_types: set[str] = set()
    for item in EQUIPMENT + NON_AIR_ROWS:
        type_name = item["type"]
        if type_name not in seen_types:
            seen_types.add(type_name)
            type_rows.append(
                [
                    type_name,
                    "Synth Schedule Generator",
                    FIXED_STAMP,
                    "Equipment",
                    "",
                    "",
                    item["capacity"],
                ]
            )
        tag = item["tag"]
        component_rows.append(
            _row(
                _component_name(type_name, tag),
                type_name,
                item["space"],
                item["capacity"],
                "n/a",
            )
        )
        for attr_name, value, unit in item.get("airflow", ()):
            attribute_rows.append(
                [
                    attr_name,
                    "Synth Schedule Generator",
                    FIXED_STAMP,
                    "Requirement",
                    "Component",
                    _component_name(type_name, tag),
                    value,
                    unit,
                ]
            )

    # The untagged row: stated, and honestly untagged.
    untagged = UNTAGGED_ROW
    if untagged["type"] not in seen_types:
        seen_types.add(untagged["type"])
        type_rows.append(
            [
                untagged["type"],
                "Synth Schedule Generator",
                FIXED_STAMP,
                "Equipment",
                "",
                "",
                untagged["capacity"],
            ]
        )
    component_rows.append(
        _row(
            _component_name(untagged["type"], ""),
            untagged["type"],
            untagged["space"],
            untagged["capacity"],
            "n/a",
        )
    )

    # A stated order is part of the workbook: the generator must not
    # shuffle rows, so the frozen truth's row numbers stay stable.
    # The only seeded variation is a deterministic rotation of the
    # component block, which keeps the sheet from being in catalogue
    # order (a real schedule is not) without using wall-clock state.
    if component_rows:
        offset = rng.randrange(len(component_rows))
        component_rows = component_rows[offset:] + component_rows[:offset]
    return component_rows, type_rows, space_rows, attribute_rows


def _workbook_bytes(component_rows, type_rows, space_rows, attribute_rows) -> bytes:
    """The workbook as deterministic bytes (same inputs, same sha256)."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.properties.creator = "scripts/make_synthetic_schedule_workbook.py"
    wb.properties.lastModifiedBy = "scripts/make_synthetic_schedule_workbook.py"
    wb.properties.created = FIXED_STAMP
    wb.properties.modified = FIXED_STAMP
    wb.properties.title = "Synthetic office schedule"

    comp = wb.active
    comp.title = "Component"
    comp.append(COMPONENT_HEADER)
    for row in component_rows:
        comp.append(row)

    ty = wb.create_sheet("Type")
    ty.append(TYPE_HEADER)
    for row in type_rows:
        ty.append(row)

    sp = wb.create_sheet("Space")
    sp.append(SPACE_HEADER)
    for row in space_rows:
        sp.append(row)

    at = wb.create_sheet("Attribute")
    at.append(ATTRIBUTE_HEADER)
    for row in attribute_rows:
        at.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return _canonical_zip(buf.getvalue())


def _canonical_zip(raw: bytes) -> bytes:
    """Rewrite the xlsx container with fixed entry timestamps.

    openpyxl stores each zip member with the current time, so two runs
    of the generator would produce different bytes (and different
    sha256) for the same content. Fixing the member timestamps makes the
    workbook byte-identical across runs, which is what pins the frozen
    truth to the workbook.
    """
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                new = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_DATE)
                new.compress_type = zipfile.ZIP_DEFLATED
                new.external_attr = 0o600 << 16
                data = src.read(info.filename)
                if info.filename == "docProps/core.xml":
                    data = _fix_core_modified(data)
                dst.writestr(new, data)
    return out.getvalue()


_MODIFIED_RE = re.compile(
    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)"
)


def _fix_core_modified(data: bytes) -> bytes:
    """Pin the workbook's modified stamp to the fixed one.

    openpyxl stamps ``docProps/core.xml`` with the wall clock when it
    saves, so two runs of the generator would differ in that one field.
    The stamp carries no truth, so it is fixed here.
    """
    stamp = FIXED_STAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    return _MODIFIED_RE.sub(rb"\g<1>" + stamp + rb"\g<2>", data)


def write_workbook(out: Path, seed: int = SEED, force: bool = False) -> Path:
    """Write the synthetic workbook to ``out``; refuse to clobber silently."""
    out = Path(out).expanduser()
    if out.exists() and not force:
        raise FileExistsError(f"{out} exists; pass --force to overwrite")
    out.parent.mkdir(parents=True, exist_ok=True)
    rows = build_rows(seed)
    out.write_bytes(_workbook_bytes(*rows))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--out", required=True, help="workbook path to write")
    ap.add_argument("--seed", type=int, default=SEED, help="generator seed")
    ap.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    args = ap.parse_args(argv)
    try:
        out = write_workbook(Path(args.out), seed=args.seed, force=args.force)
    except FileExistsError as e:
        print(str(e), file=sys.stderr)
        return 1
    import hashlib

    digest = hashlib.sha256(out.read_bytes()).hexdigest()
    print(f"wrote {out} (seed {args.seed}, sha256 {digest})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
